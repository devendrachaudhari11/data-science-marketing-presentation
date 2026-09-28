from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

# Create presentation
prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)

# Define colors
BLUE = RGBColor(0, 51, 102)  # Dark Blue
LIGHT_BLUE = RGBColor(0, 102, 204)  # Light Blue
YELLOW = RGBColor(255, 223, 0)  # Yellow
WHITE = RGBColor(255, 255, 255)
DARK_GRAY = RGBColor(51, 51, 51)

def add_title_slide(prs, title, subtitle):
    """Add a title slide"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = BLUE
    
    # Add title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2), Inches(9), Inches(2))
    title_frame = title_box.text_frame
    title_frame.word_wrap = True
    p = title_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(60)
    p.font.bold = True
    p.font.color.rgb = YELLOW
    p.alignment = PP_ALIGN.CENTER
    
    # Add subtitle
    subtitle_box = slide.shapes.add_textbox(Inches(0.5), Inches(4.2), Inches(9), Inches(2))
    subtitle_frame = subtitle_box.text_frame
    subtitle_frame.word_wrap = True
    p = subtitle_frame.paragraphs[0]
    p.text = subtitle
    p.font.size = Pt(28)
    p.font.color.rgb = YELLOW
    p.alignment = PP_ALIGN.CENTER

def add_content_slide(prs, title, content_points):
    """Add a content slide with bullet points"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = WHITE
    
    # Add blue header bar
    header_shape = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(1))
    header_shape.fill.solid()
    header_shape.fill.fore_color.rgb = BLUE
    header_shape.line.color.rgb = BLUE
    
    # Add title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.15), Inches(9), Inches(0.7))
    title_frame = title_box.text_frame
    p = title_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = YELLOW
    
    # Add content
    content_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(8.4), Inches(5.7))
    text_frame = content_box.text_frame
    text_frame.word_wrap = True
    
    for i, point in enumerate(content_points):
        if i == 0:
            p = text_frame.paragraphs[0]
        else:
            p = text_frame.add_paragraph()
        p.text = point
        p.font.size = Pt(18)
        p.font.color.rgb = DARK_GRAY
        p.space_before = Pt(6)
        p.space_after = Pt(6)
        p.level = 0

# Slide 1: Title Slide
add_title_slide(prs, "DATA SCIENCE IN MARKETING", "Comprehensive Guide to 20 Essential Topics")

# Slide 2: Introduction with Team Members
slide = prs.slides.add_slide(prs.slide_layouts[6])
background = slide.background
fill = background.fill
fill.solid()
fill.fore_color.rgb = BLUE

title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(0.8))
title_frame = title_box.text_frame
p = title_frame.paragraphs[0]
p.text = "INTRODUCTION"
p.font.size = Pt(48)
p.font.bold = True
p.font.color.rgb = YELLOW
p.alignment = PP_ALIGN.CENTER

intro_text = """Data science in marketing uses statistics, machine learning, artificial intelligence, and data visualization to understand customers, predict behavior, improve campaigns, and increase revenue.

Presented by:"""

intro_box = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(8), Inches(1.8))
intro_frame = intro_box.text_frame
intro_frame.word_wrap = True
p = intro_frame.paragraphs[0]
p.text = intro_text
p.font.size = Pt(18)
p.font.color.rgb = YELLOW

# Team members
team_members = [
    "1. Devendra Chaudhari",
    "2. Rajesh Kumar",
    "3. Priya Sharma",
    "4. Arjun Patel",
    "5. Neha Singh"
]

members_box = slide.shapes.add_textbox(Inches(2), Inches(3.5), Inches(6), Inches(3.5))
members_frame = members_box.text_frame
members_frame.word_wrap = True

for i, member in enumerate(team_members):
    if i == 0:
        p = members_frame.paragraphs[0]
    else:
        p = members_frame.add_paragraph()
    p.text = member
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = YELLOW
    p.space_after = Pt(8)

# Slide 3-22: Content Slides for 20 Topics
topics = [
    {
        "title": "1. Customer Segmentation",
        "points": [
            "• Dividing customers into groups based on shared characteristics",
            "• Common variables: Age, gender, location, income, purchase history",
            "• Methods: Demographic, geographic, behavioral, psychographic segmentation",
            "• RFM analysis: Recency, frequency, monetary value",
            "• Clustering algorithms: K-Means, hierarchical clustering, DBSCAN",
            "• Benefits: Personalized campaigns, improved satisfaction, reduced waste"
        ]
    },
    {
        "title": "2. Customer Lifetime Value Prediction",
        "points": [
            "• CLV = Average Purchase Value × Frequency × Customer Lifespan",
            "• Estimates total revenue from a customer during entire relationship",
            "• Important data: Order value, purchase frequency, retention rate, margins",
            "• Methods: Regression models, survival analysis, neural networks",
            "• Helps identify valuable customers and optimize acquisition spending",
            "• Enables targeted retention strategies for high-value customers"
        ]
    },
    {
        "title": "3. Customer Churn Prediction",
        "points": [
            "• Identifies customers likely to stop buying or cancel subscriptions",
            "• Critical for subscription businesses: streaming, telecom, software",
            "• Churn indicators: Reduced login frequency, fewer purchases, complaints",
            "• Techniques: Logistic regression, random forests, gradient boosting",
            "• Important metric: Recall (missing at-risk customers is costly)",
            "• Challenge: Determining which retention action will work best"
        ]
    },
    {
        "title": "4. Marketing Campaign Analysis",
        "points": [
            "• Measures effectiveness of advertising and promotional campaigns",
            "• Key metrics: Impressions, reach, CTR, conversion rate, ROAS",
            "• Data sources: Ad platforms, analytics, CRM, email tools, social media",
            "• Analysis includes: Cohort comparison, attribution modeling, A/B testing",
            "• Advanced: Incrementality testing, predictive optimization",
            "• Focus on final outcomes (revenue, profit, leads) not just reach/clicks"
        ]
    },
    {
        "title": "5. Market Basket Analysis",
        "points": [
            "• Studies which products customers purchase together",
            "• Key concepts: Support (frequency), Confidence, Lift",
            "• Lift = P(A and B) / [P(A) × P(B)]  — Lift > 1 indicates positive association",
            "• Algorithms: Apriori, FP-Growth, association rule mining",
            "• Applications: Cross-selling, product recommendations, bundling",
            "• Example: Customers buying cameras often buy memory cards and bags"
        ]
    },
    {
        "title": "6. Recommendation Systems",
        "points": [
            "• Suggests products or content customers may find useful",
            "• Types: Collaborative filtering, content-based, hybrid, knowledge-based",
            "• Data: Purchase history, ratings, browsing behavior, time spent",
            "• Challenges: Cold-start problem, new users, biased recommendations",
            "• Evaluation metrics: Precision, recall, click-through rate, revenue impact",
            "• Balances relevance, diversity, novelty, and business objectives"
        ]
    },
    {
        "title": "7. Predictive Analytics",
        "points": [
            "• Uses historical data and ML to predict future outcomes",
            "• Predicts: Customer purchases, lead conversion, churn, demand, sales",
            "• Common techniques: Regression, decision trees, gradient boosting, ARIMA",
            "• Process: Define problem → Collect data → Clean → Build model → Deploy",
            "• Difference: Prediction (accuracy) vs. Explanation (interpretability)",
            "• Requires regular updates as customer behavior changes"
        ]
    },
    {
        "title": "8. Marketing Attribution Modeling",
        "points": [
            "• Determines how different channels contribute to conversions",
            "• Models: First-touch, last-touch, linear, time-decay, position-based",
            "• Data-driven attribution uses statistical/ML models for credit allocation",
            "• Benefits: Understand customer journey, allocate budgets, optimize channels",
            "• Challenges: Multi-device tracking, offline interactions, cookie blocking",
            "• Combine with incrementality testing for better decision-making"
        ]
    },
    {
        "title": "9. Lead Scoring",
        "points": [
            "• Ranks potential customers by likelihood of conversion",
            "• Types: Rule-based (manual points) and Predictive (ML-based)",
            "• Data: Company size, job position, website activity, email engagement",
            "• Benefits: Prioritize leads, reduce response time, improve productivity",
            "• Evaluation: Based on actual conversion, revenue, and sales quality",
            "• Challenge: System becomes inaccurate if behavior changes; needs review"
        ]
    },
    {
        "title": "10. Conversion Rate Optimization",
        "points": [
            "• Improves percentage of visitors completing a desired action",
            "�� Conversion = (Completed Actions / Visitors) × 100",
            "• Areas analyzed: Landing pages, forms, checkout, CTAs, load speed",
            "• Methods: Funnel analysis, A/B testing, heatmaps, session analysis",
            "• Example: Testing short vs. long checkout forms to improve purchases",
            "• Focus on business outcomes (sales, leads) not vanity metrics (clicks)"
        ]
    },
    {
        "title": "11. Customer Journey Analysis",
        "points": [
            "• Studies steps customers take before, during, and after purchase",
            "• Stages: Awareness → Consideration → Evaluation → Purchase → Advocacy",
            "• Data sources: Website, social media, email, CRM, support, app usage",
            "• Methods: Funnel analysis, path analysis, process mining, Markov chains",
            "• Identifies: Drop-off points, frustrations, personalization opportunities",
            "• Journey is non-linear across web, apps, stores, and social platforms"
        ]
    },
    {
        "title": "12. Sentiment Analysis",
        "points": [
            "• Uses NLP to identify emotional tone in text (positive/negative/neutral)",
            "• Advanced: Detects specific emotions (anger, happiness, frustration)",
            "• Sources: Reviews, social media, surveys, support tickets, feedback",
            "• Methods: Keyword analysis, ML classification, NLP, transformers",
            "• Aspect-based: Analyzes sentiment for individual product features",
            "• Applications: Brand monitoring, product improvement, competitor analysis"
        ]
    },
    {
        "title": "13. Social Media Analytics",
        "points": [
            "• Analyzes data from Instagram, Facebook, LinkedIn, TikTok, YouTube, X",
            "• Metrics: Followers, reach, impressions, likes, shares, engagement rate",
            "• Engagement Rate = (Likes + Comments + Shares) / Reach × 100",
            "• Applications: Audience analysis, content comparison, influencer evaluation",
            "• Challenges: Algorithm changes, limited data access, bot engagement",
            "• More useful when connected with website visits, leads, and sales data"
        ]
    },
    {
        "title": "14. SEO Analytics",
        "points": [
            "• Improves website visibility in unpaid search engine results",
            "• Metrics: Organic traffic, keyword rankings, impressions, backlinks, CTR",
            "• Core Web Vitals: Page speed and performance indicators",
            "• Applications: Keyword research, search intent classification, content gaps",
            "• Techniques: NLP, topic modeling, keyword clustering, competitor analysis",
            "• Evaluated by: Qualified traffic, leads, sales (not just rankings)"
        ]
    },
    {
        "title": "15. Search Engine Marketing Analytics",
        "points": [
            "• Measures performance of paid search engine advertising (SEM)",
            "• Metrics: Impressions, clicks, CTR, cost per click, quality score, ROAS",
            "• Applications: Keyword analysis, bid optimization, budget allocation",
            "• Predictive bidding: ML adjusts bids based on device, location, time, user",
            "• Conversion prediction: Estimates likelihood of conversion per search",
            "• Key distinction: SEO (unpaid) vs. SEM (paid) have different strategies"
        ]
    },
    {
        "title": "16. Email Marketing Analytics",
        "points": [
            "• Measures how recipients interact with marketing emails",
            "• Metrics: Delivery rate, open rate, CTR, click-to-open rate, conversion rate",
            "• Applications: Segmentation, send-time optimization, subject line testing",
            "• Personalization: Names, recommendations, birthday offers, location content",
            "• Challenge: Open rates unreliable due to privacy features",
            "• Better metrics: Clicks, conversions, revenue, unsubscribe rates"
        ]
    },
    {
        "title": "17. A/B Testing & Experimentation",
        "points": [
            "• Compares two versions to determine which performs better",
            "• Test elements: Headlines, images, pricing, email subject lines, designs",
            "• Process: Hypothesis → Random split → Measure → Test significance → Implement",
            "• Key concepts: Control group, sample size, statistical power, confidence",
            "• Common mistakes: Ending early, too many changes, ignoring seasonality",
            "• Stronger evidence than comparing historical results"
        ]
    },
    {
        "title": "18. Sales Forecasting",
        "points": [
            "• Predicts future sales, revenue, orders, or demand",
            "• Used for: Inventory planning, staffing, budgeting, production, cash-flow",
            "• Data: Historical sales, seasonality, promotions, holidays, competition",
            "• Methods: ARIMA, exponential smoothing, Prophet, random forests, neural nets",
            "• Metrics: MAE, MSE, RMSE, MAPE, forecast bias",
            "• Include uncertainty ranges; unexpected events can make forecasts inaccurate"
        ]
    },
    {
        "title": "19. Personalization & Targeted Advertising",
        "points": [
            "• Delivers content/offers based on individual characteristics and behavior",
            "• Examples: Customized content, recommendations, location-based offers",
            "• Data: Browsing history, purchases, demographics, device, location, context",
            "• Methods: Recommendation systems, segmentation, propensity modeling",
            "• Benefits: Higher engagement, better experience, improved conversion rates",
            "• Risks: Privacy concerns, discriminatory targeting; must use clear consent"
        ]
    },
    {
        "title": "20. Fraud Detection in Digital Marketing",
        "points": [
            "• Identifies fake, deceptive, or invalid activities causing financial loss",
            "• Types: Click fraud, fake impressions, bot traffic, account takeover",
            "• Signals: High click frequency, short sessions, suspicious IPs, impossible speeds",
            "• Methods: Anomaly detection, classification, device fingerprinting, networks",
            "• Metrics: False positive/negative rates, fraud detection rate, cost prevented",
            "• Challenge: Fraudsters evolve methods; combine automated + human review"
        ]
    }
]

for topic in topics:
    add_content_slide(prs, topic["title"], topic["points"])

# Slide 23: Thank You Slide
slide = prs.slides.add_slide(prs.slide_layouts[6])
background = slide.background
fill = background.fill
fill.solid()
fill.fore_color.rgb = BLUE

thank_you_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(9), Inches(1.5))
thank_you_frame = thank_you_box.text_frame
p = thank_you_frame.paragraphs[0]
p.text = "THANK YOU"
p.font.size = Pt(72)
p.font.bold = True
p.font.color.rgb = YELLOW
p.alignment = PP_ALIGN.CENTER

thank_you_text = slide.shapes.add_textbox(Inches(0.5), Inches(4.2), Inches(9), Inches(2))
thank_you_text_frame = thank_you_text.text_frame
p = thank_you_text_frame.paragraphs[0]
p.text = "Data Science in Marketing\n\nFor Questions and Feedback, Please Contact Us"
p.font.size = Pt(24)
p.font.color.rgb = YELLOW
p.alignment = PP_ALIGN.CENTER

# Save presentation
prs.save('Data_Science_in_Marketing.pptx')
print("PowerPoint presentation created successfully: Data_Science_in_Marketing.pptx")
